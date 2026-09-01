"""Relatorio final auditavel do pipeline hibrido S10-B."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


COLUMNS = [
    "pipeline_id",
    "lote_id",
    "status_decisao",
    "regras_violadas",
    "observacao",
    "causa_provavel",
    "origem_decisao",
    "confianca_ml",
    "motivo_fallback",
    "latencia_ml_ms",
    "produto",
    "status_original",
    "status_normalizado",
    "avisos",
    "estoque_encontrado",
    "produto_estoque",
    "quantidade_estoque",
    "quantidade_pedido",
    "desktop_disponivel",
    "web_disponivel",
]


def generate_pipeline_report(results: list[dict], output_path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    result_sheet = workbook.create_sheet("Resultados")

    statuses = Counter(result.get("status_decisao") for result in results)
    divergences = [r for r in results if r.get("status_decisao") == "DIVERGENCIA"]
    fallbacks = [r for r in divergences if r.get("origem_decisao") == "fallback"]

    summary_rows = [
        ("Indicador", "Valor"),
        ("Total processado", len(results)),
        ("Validos", statuses.get("VALIDO", 0)),
        ("Divergencias", statuses.get("DIVERGENCIA", 0)),
        ("Pendentes de revisao", statuses.get("PENDENTE_REVISAO", 0)),
        ("Divergencias com fallback", len(fallbacks)),
        (
            "Pipeline operando sem ML",
            "SIM" if divergences and len(fallbacks) == len(divergences) else "NAO",
        ),
        (
            "Pipeline degradado",
            "SIM"
            if any(
                result.get("desktop_disponivel") is False
                or result.get("web_disponivel") is False
                for result in results
            )
            else "NAO",
        ),
    ]
    for row in summary_rows:
        summary_sheet.append(row)

    result_sheet.append(COLUMNS)
    for result in results:
        result_sheet.append(
            [
                _serialize_value(result.get(column))
                for column in COLUMNS
            ]
        )

    for sheet in (summary_sheet, result_sheet):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            width = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=10,
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = min(width + 2, 60)

    workbook.save(output_path)
    workbook.close()
    return output_path


def _serialize_value(value):
    if isinstance(value, (list, tuple, set)):
        return " | ".join(str(item) for item in value)
    return value
