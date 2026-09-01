"""Coletores do sistema desktop simulado e do portal web de fornecedores."""

from __future__ import annotations

import json
import subprocess
import sys
import time
from pathlib import Path

import openpyxl

from config import (
    ARQUIVO_INSPECAO,
    BASE_DIR,
    CAPSTONE_DESKTOP_EXPORT_TIMEOUT_SECONDS,
    CAPSTONE_DESKTOP_MODE,
    CAPSTONE_DESKTOP_STARTUP_SECONDS,
    CAPSTONE_WEB_MODE,
    LOGS_DIR,
)
from src.capstone_errors import InfrastructureError


def load_inventory_from_reference(path=ARQUIVO_INSPECAO) -> list[dict]:
    """Le a exportacao que representa o cadastro do sistema desktop."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Base_Referencia" not in workbook.sheetnames:
            raise ValueError("Aba Base_Referencia nao encontrada")
        sheet = workbook["Base_Referencia"]
        headers = [cell.value for cell in sheet[2]]
        records = []
        for values in sheet.iter_rows(min_row=3, values_only=True):
            if not values or all(value is None for value in values):
                continue
            raw = dict(zip(headers, values))
            if not raw.get("lote_id") or not raw.get("codigo_produto"):
                continue
            records.append(
                {
                    "lote_id": raw.get("lote_id"),
                    "produto": raw.get("codigo_produto"),
                    "descricao_produto": raw.get("descricao_produto"),
                    "status_estoque": raw.get("status_cadastro"),
                    "quantidade_estoque": 1,
                }
            )
        return records
    finally:
        workbook.close()


def load_orders_from_inspection(path=ARQUIVO_INSPECAO) -> list[dict]:
    """Le os pedidos/lotes exibidos pelo portal web simulado."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(
            (
                candidate
                for candidate in workbook.worksheets
                if candidate.title not in {"Base_Referencia", "Formulario_Analise"}
            ),
            None,
        )
        if sheet is None:
            raise ValueError("Aba de inspecao nao encontrada")
        header_row = None
        for row_number in range(1, min(sheet.max_row, 20) + 1):
            values = [sheet.cell(row_number, column).value for column in range(1, 9)]
            if values and values[0] == "lote_id":
                header_row = row_number
                headers = values
                break
        if header_row is None:
            raise ValueError("Cabecalho da inspecao nao encontrado")

        records = []
        for values in sheet.iter_rows(min_row=header_row + 1, max_col=8, values_only=True):
            if not values or all(value is None for value in values):
                continue
            if str(values[0] or "").strip().lower().startswith("total de registros"):
                break
            populated = sum(
                value is not None and bool(str(value).strip()) for value in values
            )
            if populated < 4:
                continue
            record = dict(zip(headers, values))
            record["quantidade_pedido"] = 1
            records.append(record)
        return records
    finally:
        workbook.close()


class DesktopInventoryCollector:
    """Coleta via UI real quando solicitado e usa simulacao deterministica no CI."""

    def __init__(
        self,
        *,
        mode: str = CAPSTONE_DESKTOP_MODE,
        source_path=ARQUIVO_INSPECAO,
        logs_dir=LOGS_DIR,
        startup_seconds: float = CAPSTONE_DESKTOP_STARTUP_SECONDS,
        export_timeout_seconds: float = CAPSTONE_DESKTOP_EXPORT_TIMEOUT_SECONDS,
        sleep_fn=time.sleep,
        desktop_bot_factory=None,
    ):
        self.mode = mode
        self.source_path = Path(source_path)
        self.logs_dir = Path(logs_dir)
        self.startup_seconds = startup_seconds
        self.export_timeout_seconds = export_timeout_seconds
        self.sleep_fn = sleep_fn
        self.desktop_bot_factory = desktop_bot_factory

    def collect(self, pipeline_id: str) -> dict:
        if self.mode == "simulated":
            records = load_inventory_from_reference(self.source_path)
            return {
                "status": "sucesso",
                "mode": self.mode,
                "records": records,
                "screenshot": None,
            }
        if self.mode != "visual":
            raise ValueError("CAPSTONE_DESKTOP_MODE deve ser simulated ou visual")
        return self._collect_visual(pipeline_id)

    def _collect_visual(self, pipeline_id: str) -> dict:
        self.logs_dir.mkdir(parents=True, exist_ok=True)
        safe_id = _safe_id(pipeline_id)
        export_path = self.logs_dir / f"pipeline_{safe_id}_desktop_export.json"
        screenshot_path = self.logs_dir / f"pipeline_{safe_id}_desktop.png"
        command = [
            sys.executable,
            str(BASE_DIR / "desktop_app.py"),
            "--source",
            str(self.source_path),
            "--export",
            str(export_path),
        ]
        process = subprocess.Popen(command)
        try:
            self.sleep_fn(self.startup_seconds)
            bot = self._build_desktop_bot()
            first_lot = load_inventory_from_reference(self.source_path)[0]["lote_id"]
            bot.type_key(str(first_lot))
            bot.enter()
            bot.type_keys(["ctrl", "e"])
            _wait_for_file(
                export_path,
                timeout_seconds=self.export_timeout_seconds,
                sleep_fn=self.sleep_fn,
            )
            bot.save_screenshot(str(screenshot_path))
            records = json.loads(export_path.read_text(encoding="utf-8"))
            return {
                "status": "sucesso",
                "mode": self.mode,
                "records": records,
                "screenshot": str(screenshot_path),
            }
        finally:
            process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()

    def _build_desktop_bot(self):
        if self.desktop_bot_factory is not None:
            return self.desktop_bot_factory()
        try:
            from botcity.core import DesktopBot
        except ImportError as error:
            raise InfrastructureError(
                "botcity-framework-core nao esta instalado"
            ) from error
        return DesktopBot()


class WebOrderCollector:
    def __init__(self, *, mode: str = CAPSTONE_WEB_MODE, source_path=ARQUIVO_INSPECAO):
        self.mode = mode
        self.source_path = Path(source_path)

    def collect(self, pipeline_id: str) -> dict:
        records = load_orders_from_inspection(self.source_path)
        if self.mode == "simulated":
            return {
                "status": "sucesso",
                "mode": self.mode,
                "records": records,
                "evidences": [],
            }
        if self.mode != "visual":
            raise ValueError("CAPSTONE_WEB_MODE deve ser simulated ou visual")

        from src.web_automation import processar_datapool

        automation_result = processar_datapool(return_evidencias=True)
        return {
            "status": "sucesso",
            "mode": self.mode,
            "records": records,
            "evidences": list(automation_result.get("evidencias") or []),
        }


def _wait_for_file(path: Path, *, timeout_seconds: float, sleep_fn=time.sleep):
    started_at = time.monotonic()
    while not path.exists():
        if time.monotonic() - started_at >= timeout_seconds:
            raise InfrastructureError(f"Timeout aguardando exportacao desktop: {path}")
        sleep_fn(min(0.25, timeout_seconds))


def _safe_id(value: str) -> str:
    return "".join(char if char.isalnum() or char in "-_" else "_" for char in value)
